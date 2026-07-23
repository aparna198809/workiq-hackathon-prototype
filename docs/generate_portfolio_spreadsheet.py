"""Generate Vendor Contract Portfolio Summary spreadsheet (FILE-005)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os

wb = Workbook()

BLUE_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
LIGHT_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
AMBER_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
WHITE_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
BOLD_FONT = Font(name="Calibri", bold=True, size=10)
NORMAL_FONT = Font(name="Calibri", size=10)
TITLE_FONT = Font(name="Calibri", bold=True, size=14, color="1F4E79")
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin")
)
WRAP = Alignment(wrap_text=True, vertical="top")


def style_header(ws, row, col_count):
    for c in range(1, col_count + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = WHITE_FONT
        cell.fill = BLUE_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER


def style_data(ws, start_row, end_row, col_count, sla_col=None):
    for r in range(start_row, end_row + 1):
        for c in range(1, col_count + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = NORMAL_FONT
            cell.alignment = WRAP
            cell.border = THIN_BORDER
            if (r - start_row) % 2 == 0:
                cell.fill = LIGHT_FILL
        if sla_col:
            val = ws.cell(row=r, column=sla_col).value
            if val:
                if "Green" in str(val):
                    ws.cell(row=r, column=sla_col).fill = GREEN_FILL
                elif "Amber" in str(val):
                    ws.cell(row=r, column=sla_col).fill = AMBER_FILL


def set_col_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ═══════════════════════════════════════════════════════════
# Sheet 1: Contract Overview
# ═══════════════════════════════════════════════════════════
ws1 = wb.active
ws1.title = "Contract Overview"
ws1.merge_cells("A1:K1")
ws1["A1"].value = "Vendor Contract Portfolio Summary — Q3 2026"
ws1["A1"].font = TITLE_FONT
ws1.merge_cells("A2:K2")
ws1["A2"].value = "Northbridge Health Network  |  Prepared by: Karen Liu (VMO)  |  Date: 11-JUL-2026"
ws1["A2"].font = Font(name="Calibri", italic=True, size=10, color="666666")

headers = ["VND ID", "Vendor", "Category", "Contract Type", "Value",
           "SLA Status", "SLA Response", "SLA Uptime", "Service Credit Clause",
           "Renewal Date", "Renewal Type"]
for c, h in enumerate(headers, 1):
    ws1.cell(row=4, column=c, value=h)
style_header(ws1, 4, len(headers))

data = [
    ["VND-001", "Lumina Health Systems", "EHR & Clinical Software",
     "Implementation & Support MSA", "$1,200,000", "Green",
     "4h P1, 8h P2", "99.5%",
     "MSA clause 7.4 — 8% of impl. fee per week of delay",
     "01-MAR-2027", "Fixed term — 120-day notice"],
    ["VND-002", "MediTech Biomedical Inc.", "Medical & Biomedical Equipment",
     "PM & Calibration MSA", "$340,000", "Green",
     "48h preventive, 4h critical", "N/A",
     "10% of monthly fee per month with SLA breach",
     "01-SEP-2026", "Auto-renewal — 60-day notice"],
    ["VND-003", "PharmaLink Distributors", "Clinical Supplies & Pharmacy",
     "Distribution Agreement", "$520,000", "Amber",
     "5 biz days standard, 24h urgent", "N/A",
     "Clause 5.2 — 5% line-item value per week of delay",
     "15-OCT-2026", "Annual — 90-day rebid notice"],
    ["VND-004", "ClearPoint Clinical Solutions", "Clinical Software (CPOE & Lab)",
     "License & Support Agreement", "$280,000", "Green",
     "2h P1, 8h P2", "99.5%",
     "5% monthly license per month below 99.5% uptime",
     "01-MAR-2027", "Fixed term — 120-day notice"],
]
for r, row in enumerate(data, 5):
    for c, val in enumerate(row, 1):
        ws1.cell(row=r, column=c, value=val)
style_data(ws1, 5, 8, len(headers), sla_col=6)
set_col_widths(ws1, [10, 28, 30, 28, 14, 12, 24, 10, 40, 14, 28])

# ═══════════════════════════════════════════════════════════
# Sheet 2: Performance KPIs
# ═══════════════════════════════════════════════════════════
ws2 = wb.create_sheet("Performance KPIs")
ws2.merge_cells("A1:H1")
ws2["A1"].value = "Performance KPIs by Vendor"
ws2["A1"].font = TITLE_FONT

headers2 = ["VND ID", "Vendor", "KPI 1", "Target 1", "KPI 2", "Target 2", "KPI 3", "Target 3"]
for c, h in enumerate(headers2, 1):
    ws2.cell(row=3, column=c, value=h)
style_header(ws2, 3, len(headers2))

kpi_data = [
    ["VND-001", "Lumina Health Systems",
     "Go-live on schedule", "Met ✓",
     "Post-go-live incidents", "< 5/month",
     "Training completion", "≥ 95%"],
    ["VND-002", "MediTech Biomedical Inc.",
     "Calibration completion rate", "≥ 98%",
     "Certificate turnaround", "≤ 48h",
     "Missed JC equipment items", "Zero"],
    ["VND-003", "PharmaLink Distributors",
     "On-time delivery rate", "≥ 95%",
     "Formulary-kit fill rate", "≥ 99%",
     "Controlled-substance compliance", "100%"],
    ["VND-004", "ClearPoint Clinical Solutions",
     "System uptime", "≥ 99.5%",
     "P1 incident response", "≤ 2h",
     "Upgrade deployment", "On schedule"],
]
for r, row in enumerate(kpi_data, 4):
    for c, val in enumerate(row, 1):
        ws2.cell(row=r, column=c, value=val)
style_data(ws2, 4, 7, len(headers2))
set_col_widths(ws2, [10, 28, 26, 14, 26, 14, 28, 14])

# ═══════════════════════════════════════════════════════════
# Sheet 3: Open Commitments & Milestones
# ═══════════════════════════════════════════════════════════
ws3 = wb.create_sheet("Commitments & Milestones")
ws3.merge_cells("A1:G1")
ws3["A1"].value = "Open Commitments & Upcoming Milestones"
ws3["A1"].font = TITLE_FONT

headers3 = ["VND ID", "Vendor", "Open", "Next Milestone", "Due Date", "Owner", "Status"]
for c, h in enumerate(headers3, 1):
    ws3.cell(row=3, column=c, value=h)
style_header(ws3, 3, len(headers3))

commit_data = [
    ["VND-001", "Lumina Health Systems", 0,
     "Post-go-live support period ends", "22-JUL-2026", "Karen Liu", "On Track"],
    ["VND-002", "MediTech Biomedical Inc.", 2,
     "Patient-monitor calibration (12 units)", "14-JUL-2026", "Karen Liu", "On Track"],
    ["VND-002", "MediTech Biomedical Inc.", "",
     "Autoclave validation", "18-JUL-2026", "Karen Liu", "On Track"],
    ["VND-003", "PharmaLink Distributors", 1,
     "Controlled-substance audit log binders", "16-JUL-2026", "Karen Liu", "Delayed (was 11-JUL)"],
    ["VND-004", "ClearPoint Clinical Solutions", 2,
     "CPOE v4.2 staging deployment", "17-JUL-2026", "Karen Liu", "On Track"],
    ["VND-004", "ClearPoint Clinical Solutions", "",
     "CPOE v4.2 production cutover", "21-JUL-2026", "Karen Liu", "On Track"],
]
for r, row in enumerate(commit_data, 4):
    for c, val in enumerate(row, 1):
        ws3.cell(row=r, column=c, value=val)
style_data(ws3, 4, 9, len(headers3))
# Highlight delayed row
for c in range(1, 8):
    ws3.cell(row=7, column=c).fill = AMBER_FILL
set_col_widths(ws3, [10, 28, 8, 38, 14, 14, 22])

# ═══════════════════════════════════════════════════════════
# Sheet 4: JC Compliance
# ═══════════════════════════════════════════════════════════
ws4 = wb.create_sheet("JC Compliance")
ws4.merge_cells("A1:E1")
ws4["A1"].value = "Joint Commission Vendor Compliance Status"
ws4["A1"].font = TITLE_FONT

headers4 = ["VND ID", "Vendor", "Requirement 1", "Requirement 2", "JC Readiness"]
for c, h in enumerate(headers4, 1):
    ws4.cell(row=3, column=c, value=h)
style_header(ws4, 3, len(headers4))

jc_data = [
    ["VND-001", "Lumina Health Systems",
     "Medication-management data integrity", "HIPAA BAA", "Ready ✓"],
    ["VND-002", "MediTech Biomedical Inc.",
     "Environment-of-care equipment standards", "FDA biomedical device calibration",
     "Pending (certs by 20-JUL)"],
    ["VND-003", "PharmaLink Distributors",
     "DEA controlled-substance record-keeping", "Medication-management supply chain",
     "At Risk (binder delivery 16-JUL)"],
    ["VND-004", "ClearPoint Clinical Solutions",
     "Medication-management alerting (Appendix C)", "HL7 interop with LuminaEHR",
     "On Track (cutover 21-JUL)"],
]
for r, row in enumerate(jc_data, 4):
    for c, val in enumerate(row, 1):
        ws4.cell(row=r, column=c, value=val)
style_data(ws4, 4, 7, len(headers4))
# Color-code readiness column
ws4.cell(row=4, column=5).fill = GREEN_FILL
ws4.cell(row=6, column=5).fill = AMBER_FILL
set_col_widths(ws4, [10, 28, 38, 34, 30])

# ═══════════════════════════════════════════════════════════
# Sheet 5: Renewal Calendar
# ═══════════════════════════════════════════════════════════
ws5 = wb.create_sheet("Renewal Calendar")
ws5.merge_cells("A1:F1")
ws5["A1"].value = "Contract Renewal Calendar"
ws5["A1"].font = TITLE_FONT

headers5 = ["VND ID", "Vendor", "Renewal Date", "Notice Period",
            "Notice Window Opens", "Procurement Recommendation"]
for c, h in enumerate(headers5, 1):
    ws5.cell(row=3, column=c, value=h)
style_header(ws5, 3, len(headers5))

renewal_data = [
    ["VND-002", "MediTech Biomedical Inc.", "01-SEP-2026", "60 days",
     "03-JUL-2026", "Renew with tighter 24h SLA (currently 48h)"],
    ["VND-003", "PharmaLink Distributors", "15-OCT-2026", "90 days",
     "17-JUL-2026", "Rebid controlled-substance segment; keep formulary-kit line"],
    ["VND-004", "ClearPoint Clinical Solutions", "01-MAR-2027", "120 days",
     "01-NOV-2026", "Flag v4.2 performance for renewal negotiation"],
    ["VND-001", "Lumina Health Systems", "01-MAR-2027", "120 days",
     "01-NOV-2026", "No action now"],
]
for r, row in enumerate(renewal_data, 4):
    for c, val in enumerate(row, 1):
        ws5.cell(row=r, column=c, value=val)
style_data(ws5, 4, 7, len(headers5))
set_col_widths(ws5, [10, 28, 14, 14, 18, 50])

# ═══════════════════════════════════════════════════════════
# Save
# ═══════════════════════════════════════════════════════════
out = os.path.join(
    r"c:\Users\somolinasaha\workiq\workiq-hackathon-prototype\docs",
    "Vendor_Contract_Portfolio_Summary_Q3_2026.xlsx"
)
wb.save(out)
print(f"Saved → {out}")
