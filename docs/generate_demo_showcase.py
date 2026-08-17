"""Generate the Work IQ Demo Showcase document for client presentation."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from pathlib import Path

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

BLUE = RGBColor(0x1F, 0x4E, 0x79)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
RED = RGBColor(0xC6, 0x28, 0x28)
AMBER = RGBColor(0xF5, 0x7F, 0x17)
GRAY = RGBColor(0x66, 0x66, 0x66)


def _heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = BLUE


def _bold_para(label, text):
    p = doc.add_paragraph()
    r1 = p.add_run(label)
    r1.bold = True
    r1.font.color.rgb = BLUE
    p.add_run(text)


def _table(headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            t.rows[ri + 1].cells[ci].text = str(val)


# ===================================================================
# COVER
# ===================================================================
doc.add_paragraph("\n\n\n")
title = doc.add_heading("Work IQ \u2014 Intelligent Operations Agent", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = doc.add_paragraph("Client Demo Showcase")
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.runs[0].font.size = Pt(16)
sub.runs[0].font.color.rgb = BLUE
sub.runs[0].bold = True
org = doc.add_paragraph("Northbridge Health Network")
org.alignment = WD_ALIGN_PARAGRAPH.CENTER
org.runs[0].font.size = Pt(13)
org.runs[0].font.color.rgb = GRAY
doc.add_paragraph("")
doc.add_paragraph("")
info = doc.add_paragraph("Prepared: August 2026  |  Confidential")
info.alignment = WD_ALIGN_PARAGRAPH.CENTER
info.runs[0].font.italic = True
info.runs[0].font.color.rgb = GRAY
doc.add_page_break()

# ===================================================================
# PAGE 1: PERSONA HIERARCHY (tree table)
# ===================================================================
_heading("1. Organization & Persona Hierarchy")
doc.add_paragraph(
    "Northbridge Health Network manages three critical business areas. "
    "The tree below shows reporting lines and business area ownership."
)
doc.add_paragraph("")

# Visual tree using styled paragraphs with monospace font
def _tree_node(text, level=0, is_last=False, color=None, bold=False, font_size=10):
    """Add a tree-node paragraph with visual connectors."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Inches(level * 0.4)

    if level == 0:
        connector = ""
    elif is_last:
        connector = "\u2514\u2500\u2500 "
    else:
        connector = "\u251c\u2500\u2500 "

    if connector:
        rc = p.add_run(connector)
        rc.font.name = "Consolas"
        rc.font.size = Pt(font_size)
        rc.font.color.rgb = GRAY

    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(font_size)
    r.bold = bold
    if color:
        r.font.color.rgb = color
    return p

# --- Executive ---
_tree_node("\u25c6  JAMES WHITAKER \u2014 Director of Clinic Operations",
           level=0, bold=True, color=BLUE, font_size=12)
p = doc.add_paragraph()
p.paragraph_format.left_indent = Inches(0.15)
r = p.add_run("Executive oversight \u2014 all three business areas")
r.font.italic = True
r.font.color.rgb = GRAY
r.font.size = Pt(9)

doc.add_paragraph("")

# --- Direct Reports ---
# Maria
_tree_node("MARIA DELGADO \u2014 Quality Program Manager",
           level=1, color=GREEN, bold=True)
p = doc.add_paragraph()
p.paragraph_format.left_indent = Inches(0.55)
r = p.add_run("\u25b8 Business Area: Action Items & CAPA Tracker")
r.font.size = Pt(9)
r.font.color.rgb = GREEN

_tree_node("Angela Foster \u2014 Quality Analyst (CAPA tracking)",
           level=2, is_last=True)

# Sandra
_tree_node("SANDRA OKAFOR \u2014 Credentialing Manager",
           level=1, color=GREEN, bold=True)
p = doc.add_paragraph()
p.paragraph_format.left_indent = Inches(0.55)
r = p.add_run("\u25b8 Business Area: Action Items & CAPA Tracker (Credentialing)")
r.font.size = Pt(9)
r.font.color.rgb = GREEN

_tree_node("Nina Alvarez \u2014 Credentialing Specialist",
           level=2, is_last=True)

# Karen
_tree_node("KAREN LIU \u2014 Vendor/Contract Manager (VMO)",
           level=1, color=AMBER, bold=True)
p = doc.add_paragraph()
p.paragraph_format.left_indent = Inches(0.55)
r = p.add_run("\u25b8 Business Area: Vendor Contract Management ($2.34M portfolio)")
r.font.size = Pt(9)
r.font.color.rgb = AMBER

# Marcus
_tree_node("MARCUS WEBB \u2014 Procurement & Sourcing Lead",
           level=1, color=RED, bold=True)
p = doc.add_paragraph()
p.paragraph_format.left_indent = Inches(0.55)
r = p.add_run("\u25b8 Business Area: Vendor RFP & Proposal Evaluation (3 RFPs, 7 proposals)")
r.font.size = Pt(9)
r.font.color.rgb = RED

# Robert
_tree_node("ROBERT KLEIN \u2014 EHR Program Lead",
           level=1)
p = doc.add_paragraph()
p.paragraph_format.left_indent = Inches(0.55)
r = p.add_run("\u25b8 EHR Rollout (Westgate clinic)")
r.font.size = Pt(9)
r.font.color.rgb = GRAY

_tree_node("Greg Sullivan \u2014 IT Integration Engineer (contractor)",
           level=2, is_last=True)

# Priya
_tree_node("PRIYA RAMAN \u2014 Compliance Officer",
           level=1)
p = doc.add_paragraph()
p.paragraph_format.left_indent = Inches(0.55)
r = p.add_run("\u25b8 All areas \u2014 regulatory readiness, Joint Commission")
r.font.size = Pt(9)
r.font.color.rgb = GRAY

# David
_tree_node("DAVID MUNOZ \u2014 Pharmacy Ops Coordinator",
           level=1, is_last=True)
p = doc.add_paragraph()
p.paragraph_format.left_indent = Inches(0.55)
r = p.add_run("\u25b8 Quality workflows (medication reconciliation)")
r.font.size = Pt(9)
r.font.color.rgb = GRAY

# --- External Vendors Section ---
doc.add_paragraph("")
p = doc.add_paragraph()
r = p.add_run("\u2500\u2500\u2500  ONBOARDED VENDORS  ")
r.bold = True
r.font.color.rgb = AMBER
r.font.size = Pt(10)
r2 = p.add_run("(managed by Karen Liu \u2014 VMO)")
r2.font.color.rgb = GRAY
r2.font.size = Pt(9)

_tree_node("Tom Becker \u2014 Lumina Health Systems (EHR)  |  VND-001 \u2014 $1.2M MSA",
           level=1)
_tree_node("Rachel Torres \u2014 MediTech Biomedical (Equipment)  |  VND-002 \u2014 $340K MSA",
           level=1)
_tree_node("Steve Nakamura \u2014 PharmaLink Distributors (Supplies)  |  VND-003 \u2014 $520K MSA",
           level=1)
_tree_node("Lisa Chen \u2014 ClearPoint Clinical Solutions (Software)  |  VND-004 \u2014 $280K MSA",
           level=1, is_last=True)

# --- Vendor Candidates (RFP Proposals) ---
doc.add_paragraph("")
p = doc.add_paragraph()
r = p.add_run("\u2500\u2500\u2500  VENDOR CANDIDATES (RFP PROPOSALS)  ")
r.bold = True
r.font.color.rgb = RED
r.font.size = Pt(10)
r2 = p.add_run("(managed by Marcus Webb \u2014 Procurement)")
r2.font.color.rgb = GRAY
r2.font.size = Pt(9)

# RFP-001
p = doc.add_paragraph()
p.paragraph_format.left_indent = Inches(0.15)
r = p.add_run("RFP-001: Telehealth Platform")
r.bold = True
r.font.color.rgb = BLUE
r.font.size = Pt(10)

_tree_node("Dana Mitchell \u2014 VirtuCare Health  |  $185K/yr  |  Score: 7.80  \u2705 SHORTLISTED",
           level=1)
_tree_node("Ryan Patel \u2014 TeleMedix Solutions  |  $210K/yr  |  Score: 8.35  \u2705 SHORTLISTED",
           level=1)
_tree_node("Sofia Reyes \u2014 HealthBridge Connect  |  $125K/yr  |  Score: 6.45  \u274c NOT SHORTLISTED",
           level=1, is_last=True)

# RFP-002
p = doc.add_paragraph()
p.paragraph_format.left_indent = Inches(0.15)
r = p.add_run("RFP-002: Clinical Analytics & Reporting")
r.bold = True
r.font.color.rgb = BLUE
r.font.size = Pt(10)

_tree_node("Alex Thornton \u2014 Meridian Data Sciences  |  $165K/yr  |  Score: 8.10  \u2705 SHORTLISTED",
           level=1)
_tree_node("Priya Kapoor \u2014 InsightHealth Analytics  |  $130K/yr  |  Score: 6.90  \u274c NOT SHORTLISTED",
           level=1, is_last=True)

# RFP-003
p = doc.add_paragraph()
p.paragraph_format.left_indent = Inches(0.15)
r = p.add_run("RFP-003: Medical Waste Management")
r.bold = True
r.font.color.rgb = BLUE
r.font.size = Pt(10)

_tree_node("Carlos Mendez \u2014 EnviroMed Disposal  |  $82K/yr  |  Score: 7.70  \u2705 SHORTLISTED",
           level=1)
_tree_node("Laura Kim \u2014 GreenHealth Waste Services  |  $68K/yr  |  Score: 6.85  \u274c NOT SHORTLISTED",
           level=1, is_last=True)

doc.add_paragraph("")
_heading("Three Business Areas at a Glance", level=2)
_table(
    ["Business Area", "Owner", "Scope", "Key Metric"],
    [
        ["Action Items (CAPA)", "Maria Delgado", "7 corrective actions, 2 committees", "2 past-due, 2 flagged"],
        ["Vendor Contract Mgmt", "Karen Liu", "4 active vendors, $2.34M portfolio", "1 SLA amber (PharmaLink)"],
        ["Vendor RFP & Proposals", "Marcus Webb", "3 open RFPs, 7 proposals received", "4 shortlisted, 3 rejected"],
    ],
)
doc.add_page_break()

# ===================================================================
# PAGE 2: TOOLS & SCOPE
# ===================================================================
_heading("2. Tools & Data Landscape")
_heading("Tracker Systems in Scope", level=2)
_table(
    ["Tracker", "Business Area", "Current Format", "Records", "Updated By"],
    [
        ["CAPA Tracker", "Corrective Actions", "Self-managed Excel / SharePoint", "7 items", "Manual \u2014 quality team"],
        ["Vendor Contract Tracker", "Onboarded Vendor Mgmt", "Excel spreadsheet", "4 contracts ($2.34M)", "Manual \u2014 VMO team"],
        ["Vendor Proposal Tracker", "RFP & Procurement", "Does not exist", "7 proposals / 3 RFPs", "Not tracked \u2014 emails only"],
    ],
)
doc.add_paragraph("")
_heading("Communication & Collaboration Tools", level=2)
_table(
    ["Tool", "Data Type", "Monthly Volume", "Used For"],
    [
        ["Outlook Email", "Emails & threads", "1,000+", "Vendor correspondence, proposals, escalations"],
        ["Microsoft Teams", "Channel messages", "1,000+", "Program updates, vendor discussions, committee notes"],
        ["SharePoint / OneDrive", "Documents & files", "100+", "Policies, contracts, proposal PDFs, checklists"],
        ["Teams Meetings", "Recaps & action items", "20+", "Committee decisions, vendor demos, reviews"],
        ["People / Profiles", "Org directory", "22 people", "Roles, expertise, committee membership"],
    ],
)
doc.add_page_break()

# ===================================================================
# PAGE 3: THE PROBLEM
# ===================================================================
_heading("3. The Problem")

_heading("Information Overload", level=2)
p = doc.add_paragraph()
r = p.add_run("Every week, a single manager must process:")
r.bold = True
r.font.size = Pt(12)

for prob in [
    "250+ emails across vendor threads, committee updates, escalation chains, and proposals",
    "250+ Teams messages across Quality, Vendor Management, Procurement, and EHR Rollout channels",
    "Dozens of policy documents, vendor contracts, and proposal PDFs in SharePoint",
    "5+ meeting recaps with action items buried in calendar notes and recap threads",
    "No single view connecting: email thread \u2192 meeting decision \u2192 tracker update",
]:
    doc.add_paragraph(prob, style="List Bullet")

doc.add_paragraph("")
_heading("No Unified Tracking \u2014 Three Areas, Zero Automation", level=2)

for label, text in [
    ("\u274c Action Items (CAPAs): ",
     "Actions are created in meetings, discussed in emails, and manually updated in Excel. "
     "No automated detection of new action items. No reconciliation between what a meeting "
     "decided and what the tracker shows. Items go stale. Past-due items are discovered too "
     "late \u2014 putting Joint Commission readiness at risk."),
    ("\u274c Vendor Contract Management: ",
     "SLA performance, delivery delays, and renewal dates live in a separate Excel. Status "
     "updates arrive via emails and Teams but are never reflected in the tracker automatically. "
     "The weekly vendor status report takes 3\u20134 DAYS of manual preparation \u2014 collating data "
     "from 5+ sources into an Excel workbook."),
    ("\u274c Vendor Proposal Evaluation: ",
     "No system exists. Proposals arrive by email with PDF attachments. Evaluators read each "
     "one manually, compare on gut feel, and discuss in ad-hoc meetings. No structured scoring. "
     "No side-by-side comparison. No audit trail. Shortlisting takes weeks."),
]:
    p = doc.add_paragraph()
    r1 = p.add_run(label)
    r1.bold = True
    r1.font.color.rgb = RED
    p.add_run(text)

doc.add_paragraph("")
_heading("Business Impact", level=2)
_table(
    ["Pain Point", "Current State", "Impact"],
    [
        ["Action item tracking", "Manual Excel, updated weekly", "Past-due CAPAs \u2192 Joint Commission audit risk"],
        ["Weekly vendor report", "3\u20134 days to prepare manually", "Leadership gets stale data \u2192 slow decisions"],
        ["Proposal evaluation", "No system \u2014 email + gut feel", "Weeks to shortlist \u2192 procurement bottleneck"],
        ["Cross-signal visibility", "Siloed email / Teams / SharePoint", "Decisions without full context \u2192 missed risks"],
        ["Compliance readiness", "Manual checklist reconciliation", "Audit findings \u2192 remediation costs"],
    ],
)
doc.add_page_break()

# ===================================================================
# PAGE 4: THE SOLUTION
# ===================================================================
_heading("4. The Solution \u2014 Work IQ Intelligent Agents")
doc.add_paragraph(
    "Work IQ deploys three specialized AI agents backed by Azure AI Search. They continuously "
    "monitor all organizational signals \u2014 emails, Teams, meetings, and documents \u2014 and take "
    "action with human-in-the-loop approval. Nothing changes without explicit sign-off."
)
doc.add_paragraph("")

# Agent 1
_heading("\u2705 Agent 1: Reconciliation Agent", level=2)
_bold_para("Scope: ", "CAPA Tracker \u2014 Corrective Action Intelligence")
for a in [
    "Monitors every new email, Teams message, and meeting recap in real time (CDC)",
    "Extracts action items automatically using semantic understanding",
    "Matches against existing CAPA tracker using Azure AI Search",
    "Proposes updates (status change, new due date) or new entries \u2014 with source citations",
    "Human-in-the-loop: nothing is written until an authorized user clicks Approve",
    "Result: a clean, always-current tracker with a full audit trail",
]:
    doc.add_paragraph(a, style="List Bullet")
p = doc.add_paragraph()
r = p.add_run("\u23f1 Time saved: ~4 hours/week ")
r.bold = True
r.font.color.rgb = GREEN
p.add_run("\u2014 past-due items surfaced in minutes, not days.")
doc.add_paragraph("")

# Agent 2
_heading("\u2705 Agent 2: Proposal Evaluation Agent", level=2)
_bold_para("Scope: ", "Vendor RFP & Proposal Scoring")
for a in [
    "Detects new vendor proposal emails as they arrive (CDC)",
    "Extracts pricing, SLAs, compliance certs, timelines using Azure AI Search (semantic retrieval)",
    "Scores each proposal against 7 weighted criteria:",
]:
    doc.add_paragraph(a, style="List Bullet")

# Criteria sub-table
ct = doc.add_table(rows=8, cols=2)
ct.style = "Light List Accent 1"
ct.rows[0].cells[0].text = "Criterion"
ct.rows[0].cells[1].text = "Weight"
for p in ct.rows[0].cells[0].paragraphs:
    for r in p.runs:
        r.bold = True
for p in ct.rows[0].cells[1].paragraphs:
    for r in p.runs:
        r.bold = True
for i, (crit, wt) in enumerate([
    ("Cost & TCO", "20%"), ("Technical Capability", "20%"),
    ("Performance & SLA", "15%"), ("Compliance & Security", "15%"),
    ("Implementation Plan", "10%"), ("Vendor Stability", "10%"),
    ("Innovation & Roadmap", "10%"),
]):
    ct.rows[i + 1].cells[0].text = crit
    ct.rows[i + 1].cells[1].text = wt

doc.add_paragraph("")
for a in [
    "Ranks proposals per RFP \u2014 every score links back to the source email/document/meeting",
    "Re-scores automatically when new evidence arrives (follow-up emails, vendor demos)",
    "Generates side-by-side comparison with risk flags for the evaluation panel",
]:
    doc.add_paragraph(a, style="List Bullet")
p = doc.add_paragraph()
r = p.add_run("\u23f1 Time saved: Weeks \u2192 minutes. ")
r.bold = True
r.font.color.rgb = GREEN
p.add_run("Structured, auditable shortlisting replaces ad-hoc decisions.")
doc.add_paragraph("")

# Agent 3
_heading("\u2705 Agent 3: Report Generator", level=2)
_bold_para("Scope: ", "Weekly Vendor Status & Joint Commission Readiness Reports")
for a in [
    "Generates Weekly Vendor Status as a formatted Excel workbook (5 sheets: Dashboard, SLA, Commitments, Renewals, Risks)",
    "Generates Joint Commission Readiness Report with vendor compliance status",
    "Pulls live data from ALL signals \u2014 emails, meetings, Teams, AND tracker tables",
    "RBAC-aware: each persona sees only the data they're authorized to access",
]:
    doc.add_paragraph(a, style="List Bullet")
p = doc.add_paragraph()
r = p.add_run("\u23f1 Time saved: 3\u20134 days \u2192 30 seconds. ")
r.bold = True
r.font.color.rgb = GREEN
p.add_run("Leadership gets real-time data, not week-old snapshots.")
doc.add_page_break()

# ===================================================================
# PAGE 5: ARCHITECTURE (table-based diagram)
# ===================================================================
_heading("5. Technical Architecture")
_heading("Signal Flow", level=2)

arch = doc.add_table(rows=7, cols=5)
arch.alignment = WD_TABLE_ALIGNMENT.CENTER

# Row 0 — Data sources
for i, (src, vol) in enumerate([
    ("Emails", "1,000+/mo"), ("Teams", "1,000+/mo"),
    ("Meetings", "20+/mo"), ("Files", "100+"),
    ("People", "22"),
]):
    cell = arch.rows[0].cells[i]
    p = cell.paragraphs[0]
    r = p.add_run(f"{src}\n{vol}")
    r.font.size = Pt(9)
    r.bold = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Row 1 — arrow
arch.rows[1].cells[2].paragraphs[0].add_run("\u25bc").font.size = Pt(14)
arch.rows[1].cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

# Row 2 — AI Search
cell = arch.rows[2].cells[1]
cell.merge(arch.rows[2].cells[3])
p = cell.paragraphs[0]
r = p.add_run("Azure AI Search \u2014 Semantic + Keyword Hybrid Retrieval")
r.bold = True
r.font.color.rgb = BLUE
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Row 3 — arrow
arch.rows[3].cells[2].paragraphs[0].add_run("\u25bc").font.size = Pt(14)
arch.rows[3].cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

# Row 4 — Three agents
for i, (name, desc) in enumerate([
    ("Reconciliation\nAgent", "CDC \u2192 CAPA Tracker\n(with approval)"),
    (None, None),
    ("Proposal Eval\nAgent", "CDC \u2192 Proposal Tracker\n(scores + citations)"),
    (None, None),
    ("Report\nGenerator", "On-demand Excel\nworkbooks"),
]):
    if name:
        cell = arch.rows[4].cells[i]
        p = cell.paragraphs[0]
        r = p.add_run(name)
        r.bold = True
        r.font.size = Pt(9)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2 = cell.add_paragraph()
        r2 = p2.add_run(desc)
        r2.font.size = Pt(8)
        r2.font.color.rgb = GRAY
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Row 5 — arrow
arch.rows[5].cells[2].paragraphs[0].add_run("\u25bc").font.size = Pt(14)
arch.rows[5].cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

# Row 6 — Human in the loop
cell = arch.rows[6].cells[1]
cell.merge(arch.rows[6].cells[3])
p = cell.paragraphs[0]
r = p.add_run("Human-in-the-Loop \u2014 Approve / Reject")
r.bold = True
r.font.color.rgb = GREEN
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph("")
_heading("Key Technology", level=2)
_table(
    ["Component", "Technology", "Purpose"],
    [
        ["Retrieval", "Azure AI Search (vector + keyword)", "Semantic hybrid search across all M365 signals"],
        ["Scoring LLM", "Azure OpenAI (GPT-4o-mini)", "Proposal data extraction and criterion scoring"],
        ["Embeddings", "text-embedding-3-small", "Vector embeddings for semantic similarity"],
        ["Agent Framework", "Microsoft Agent Framework + A2A", "Multi-agent orchestration with tool calling"],
        ["MCP Tools", "MCP stdio (simulator)", "ask, fetch, create, update, score, generate"],
    ],
)
doc.add_page_break()

# ===================================================================
# PAGE 6: VALUE SUMMARY
# ===================================================================
_heading("6. Business Value \u2014 Before vs After")
_table(
    ["Capability", "Before Work IQ", "After Work IQ", "Value Delivered"],
    [
        ["Action item\ntracking", "Manual Excel\n2\u20133 days stale", "Auto-detected from signals\nReal-time with approval", "~4 hrs/week saved\nZero stale items"],
        ["Vendor weekly\nstatus report", "3\u20134 days manual\npreparation", "One-click generation\n30 seconds, live data", "3\u20134 days \u2192 30 seconds\nAlways current"],
        ["Proposal\nevaluation", "No system\nWeeks to shortlist", "Auto-scored, 7 criteria\nRanked with citations", "Weeks \u2192 minutes\nFull audit trail"],
        ["Cross-signal\nvisibility", "Siloed tools\nContext lost", "AI Search connects\nall signals", "Complete context\nfor every decision"],
        ["Compliance\nreadiness", "Manual checklist\nRisk of gaps", "Continuous monitoring\nProactive risk flags", "Audit-ready\nat all times"],
    ],
)

doc.add_paragraph("")
_heading("Estimated Impact", level=2)

for label, text in [
    ("Per week: ", "~30+ hours saved across quality, procurement, and vendor management teams"),
    ("Per month: ", "~120+ hours \u2014 equivalent to 3 full work weeks redirected to strategic priorities"),
    ("Risk reduction: ", "Zero stale action items, structured vendor selection with audit trail, "
     "real-time compliance visibility for Joint Commission readiness"),
]:
    p = doc.add_paragraph()
    r = p.add_run(label)
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = GREEN
    p.add_run(text)

doc.add_paragraph("")
_heading("Live Demo Scenarios", level=2)
doc.add_paragraph("During the demo, watch for these real-time interactions:")
doc.add_paragraph("")

for i, (q, result) in enumerate([
    ('"What corrective actions are stalled?"',
     "Agent searches emails + meetings + tracker, returns cited answer with past-due flags"),
    ('"Score and rank the telehealth vendor proposals"',
     "Agent reads 3 proposal emails + documents, scores on 7 criteria, presents ranked table"),
    ('"Generate the weekly vendor status report"',
     "Agent produces a 5-sheet Excel workbook in ~30 seconds from live data"),
    ('"What did the evaluation panel decide about shortlisting?"',
     "Agent synthesizes meeting recap + scoring data + emails into a cited summary"),
    ("A new vendor email arrives",
     "Agent detects it automatically, re-scores the affected proposal, updates the tracker"),
], start=1):
    p = doc.add_paragraph()
    r = p.add_run(f"{i}. ")
    r.bold = True
    r.font.color.rgb = BLUE
    r2 = p.add_run(q)
    r2.bold = True
    p2 = doc.add_paragraph(f"    \u2192 {result}")
    p2.paragraph_format.left_indent = Inches(0.5)


out = Path(__file__).resolve().parent / "WORKIQ_DEMO_SHOWCASE.docx"
doc.save(str(out))
print(f"Saved: {out}")
