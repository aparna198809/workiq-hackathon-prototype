"""Rebuild Section 1 of WORKIQ_DEMO_SHOWCASE.docx with a Visio-style hierarchy diagram using tables."""

from pathlib import Path
from copy import deepcopy

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml


DOC_PATH = Path(__file__).parent / "WORKIQ_DEMO_SHOWCASE.docx"
OUT_PATH = DOC_PATH  # overwrite in place


def set_cell_shading(cell, color_hex: str):
    """Set cell background color."""
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}" w:val="clear"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_border(cell, top="single", bottom="single", left="single", right="single", color="404040", size="8"):
    """Set cell borders."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="{top}" w:sz="{size}" w:color="{color}"/>'
        f'  <w:bottom w:val="{bottom}" w:sz="{size}" w:color="{color}"/>'
        f'  <w:left w:val="{left}" w:sz="{size}" w:color="{color}"/>'
        f'  <w:right w:val="{right}" w:sz="{size}" w:color="{color}"/>'
        f'</w:tcBorders>'
    )
    tcBorders_existing = tcPr.find(qn("w:tcBorders"))
    if tcBorders_existing is not None:
        tcPr.remove(tcBorders_existing)
    tcPr.append(tcBorders)


def make_box_cell(cell, name: str, role: str, scope: str, color: str, text_color: str = "FFFFFF"):
    """Format a cell as a persona 'box' in the org chart."""
    set_cell_shading(cell, color)
    set_cell_border(cell, color=color)
    
    # Clear existing
    for p in cell.paragraphs:
        p.clear()
    
    # Name (bold)
    p_name = cell.paragraphs[0]
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_name.add_run(name)
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(text_color)
    
    # Role
    p_role = cell.add_paragraph()
    p_role.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_role.add_run(role)
    run.font.size = Pt(7.5)
    run.font.color.rgb = RGBColor.from_string(text_color)
    run.italic = True
    
    # Scope
    if scope:
        p_scope = cell.add_paragraph()
        p_scope.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p_scope.add_run(scope)
        run.font.size = Pt(7)
        run.font.color.rgb = RGBColor.from_string(text_color)
    
    # Reduce paragraph spacing
    for p in cell.paragraphs:
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)


def make_connector_cell(cell, text: str = "│"):
    """Make a cell show a downward connector line."""
    set_cell_shading(cell, "FFFFFF")
    set_cell_border(cell, top="nil", bottom="nil", left="nil", right="nil")
    cell.paragraphs[0].clear()
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string("404040")


def make_empty_cell(cell):
    """Transparent empty cell."""
    set_cell_shading(cell, "FFFFFF")
    set_cell_border(cell, top="nil", bottom="nil", left="nil", right="nil")
    cell.paragraphs[0].clear()


def make_vendor_cell(cell, name: str, vendor: str, contract: str, color: str = "E8F5E9"):
    """Format a cell as a vendor box."""
    set_cell_shading(cell, color)
    set_cell_border(cell, color="388E3C", size="6")
    
    for p in cell.paragraphs:
        p.clear()
    
    p_name = cell.paragraphs[0]
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_name.add_run(name)
    run.bold = True
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string("1B5E20")
    
    p_vendor = cell.add_paragraph()
    p_vendor.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_vendor.add_run(vendor)
    run.font.size = Pt(7)
    run.font.color.rgb = RGBColor.from_string("2E7D32")
    
    p_contract = cell.add_paragraph()
    p_contract.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_contract.add_run(contract)
    run.font.size = Pt(7)
    run.font.color.rgb = RGBColor.from_string("555555")
    
    for p in cell.paragraphs:
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)


def make_rfp_cell(cell, name: str, company: str, details: str, color: str = "FFF3E0"):
    """Format a cell as an RFP candidate box."""
    set_cell_shading(cell, color)
    set_cell_border(cell, color="E65100", size="6")
    
    for p in cell.paragraphs:
        p.clear()
    
    p_name = cell.paragraphs[0]
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_name.add_run(name)
    run.bold = True
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string("BF360C")
    
    p_co = cell.add_paragraph()
    p_co.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_co.add_run(company)
    run.font.size = Pt(7)
    run.font.color.rgb = RGBColor.from_string("E65100")
    
    p_det = cell.add_paragraph()
    p_det.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_det.add_run(details)
    run.font.size = Pt(6.5)
    run.font.color.rgb = RGBColor.from_string("555555")
    
    for p in cell.paragraphs:
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)


def build_hierarchy(doc):
    """Remove old text hierarchy (paras 9-49) and insert Visio-style table diagram."""
    # Find the section heading and the next heading
    heading_idx = None
    next_section_idx = None
    for i, para in enumerate(doc.paragraphs):
        if para.style.name == "Heading 1" and "Organization" in para.text:
            heading_idx = i
        elif heading_idx is not None and para.style.name == "Heading 2" and "Three Business Areas" in para.text:
            next_section_idx = i
            break

    if heading_idx is None or next_section_idx is None:
        raise RuntimeError("Could not locate section boundaries")

    # Remove paragraphs between heading and "Three Business Areas" subheading
    # We'll delete from index heading_idx+1 to next_section_idx-1
    body = doc.element.body
    paras_to_remove = []
    for i in range(heading_idx + 1, next_section_idx):
        paras_to_remove.append(doc.paragraphs[i]._element)
    
    for elem in paras_to_remove:
        body.remove(elem)

    # Now insert new content after the heading
    # We need to find the heading element and insert after it
    heading_elem = doc.paragraphs[heading_idx]._element

    # Build the hierarchy as a series of elements to insert
    # We'll create the tables and paragraphs, then insert them

    # Strategy: Add content at the end of document first, then move elements
    # Actually, easier: add a temporary paragraph after the heading, then build tables
    
    # Insert intro paragraph
    intro = parse_xml(
        f'<w:p {nsdecls("w")}>'
        f'  <w:pPr><w:spacing w:after="120"/></w:pPr>'
        f'  <w:r><w:rPr><w:sz w:val="20"/></w:rPr>'
        f'    <w:t>The diagram below shows the organizational hierarchy, reporting lines, business areas, and data access levels (RBAC). Each box represents a persona with its access scope.</w:t>'
        f'  </w:r>'
        f'</w:p>'
    )
    heading_elem.addnext(intro)
    
    # === TIER 0: Executive (top) ===
    # We'll build tables programmatically and insert after the intro
    
    # Create a temporary document to build tables, then transplant elements
    tmp = Document()
    
    # --- LEVEL 0: Director ---
    tmp.add_paragraph("")  # spacer
    t0 = tmp.add_table(rows=1, cols=5)
    t0.alignment = WD_TABLE_ALIGNMENT.CENTER
    make_empty_cell(t0.cell(0, 0))
    make_empty_cell(t0.cell(0, 1))
    make_box_cell(t0.cell(0, 2), "JAMES WHITAKER", "Director of Clinic Operations",
                  "🔓 FULL ACCESS — all data", "1A237E")
    make_empty_cell(t0.cell(0, 3))
    make_empty_cell(t0.cell(0, 4))
    
    # Connector row
    tc = tmp.add_table(rows=1, cols=5)
    tc.alignment = WD_TABLE_ALIGNMENT.CENTER
    make_empty_cell(tc.cell(0, 0))
    make_empty_cell(tc.cell(0, 1))
    make_connector_cell(tc.cell(0, 2), "┃")
    make_empty_cell(tc.cell(0, 3))
    make_empty_cell(tc.cell(0, 4))
    
    # Connector spreading
    tc2 = tmp.add_table(rows=1, cols=5)
    tc2.alignment = WD_TABLE_ALIGNMENT.CENTER
    make_connector_cell(tc2.cell(0, 0), "┌────────")
    make_connector_cell(tc2.cell(0, 1), "──────┬──")
    make_connector_cell(tc2.cell(0, 2), "──────┼──")
    make_connector_cell(tc2.cell(0, 3), "──────┬──")
    make_connector_cell(tc2.cell(0, 4), "────────┐")
    
    # --- LEVEL 1: Direct Reports (5 columns) ---
    t1 = tmp.add_table(rows=1, cols=5)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    make_box_cell(t1.cell(0, 0), "MARIA DELGADO", "Quality Program Mgr",
                  "CAPA Tracker\n🔓 Quality data", "1565C0")
    make_box_cell(t1.cell(0, 1), "KAREN LIU", "Vendor/Contract Mgr (VMO)",
                  "Vendor Contracts\n🔓 All vendor + SLA\n🚫 HR credentialing", "0277BD")
    make_box_cell(t1.cell(0, 2), "SANDRA OKAFOR", "Credentialing Manager",
                  "Credentialing\n🔓 HR-sensitive data\n🚫 Commercial terms", "00695C")
    make_box_cell(t1.cell(0, 3), "MARCUS WEBB", "Procurement Lead",
                  "Vendor RFP & Proposals\n🔓 Evaluation data", "4527A0")
    make_box_cell(t1.cell(0, 4), "ROBERT KLEIN", "EHR Program Lead",
                  "EHR Rollout\n🔓 Vendor rollout signals", "283593")
    
    # Connector row
    tc3 = tmp.add_table(rows=1, cols=5)
    tc3.alignment = WD_TABLE_ALIGNMENT.CENTER
    make_connector_cell(tc3.cell(0, 0), "│")
    make_connector_cell(tc3.cell(0, 1), "│")
    make_connector_cell(tc3.cell(0, 2), "│")
    make_connector_cell(tc3.cell(0, 3), "│")
    make_connector_cell(tc3.cell(0, 4), "│")
    
    # --- LEVEL 2: Sub-reports ---
    t2 = tmp.add_table(rows=1, cols=5)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    make_box_cell(t2.cell(0, 0), "Angela Foster", "Quality Analyst",
                  "CAPA tracking", "90CAF9", "1A237E")
    make_box_cell(t2.cell(0, 1), "PRIYA RAMAN", "Compliance Officer",
                  "JC Readiness\nAll areas", "81D4FA", "01579B")
    make_box_cell(t2.cell(0, 2), "Nina Alvarez", "Credentialing Specialist",
                  "Onboarding", "80CBC4", "004D40")
    make_box_cell(t2.cell(0, 3), "DAVID MUNOZ", "Pharmacy Ops",
                  "Med reconciliation", "B39DDB", "311B92")
    make_box_cell(t2.cell(0, 4), "Greg Sullivan", "IT Integration (Contractor)",
                  "🚫 Internal strategy\n🚫 Commercial terms", "BBDEFB", "1A237E")
    
    # Spacer + Vendor Section Header
    sp = tmp.add_paragraph()
    sp.paragraph_format.space_before = Pt(16)
    sp.paragraph_format.space_after = Pt(4)
    
    p_vendor_hdr = tmp.add_paragraph()
    p_vendor_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_vendor_hdr.add_run("── ONBOARDED VENDORS (managed by Karen Liu — VMO) ──")
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string("2E7D32")
    
    # --- Vendor Row ---
    tv = tmp.add_table(rows=1, cols=4)
    tv.alignment = WD_TABLE_ALIGNMENT.CENTER
    make_vendor_cell(tv.cell(0, 0), "Tom Becker", "Lumina Health Systems (EHR)", "VND-001 — $1.2M MSA")
    make_vendor_cell(tv.cell(0, 1), "Rachel Torres", "MediTech Biomedical (Equip.)", "VND-002 — $340K MSA")
    make_vendor_cell(tv.cell(0, 2), "Steve Nakamura", "PharmaLink Distributors", "VND-003 — $520K MSA  ⚠️ AMBER")
    make_vendor_cell(tv.cell(0, 3), "Lisa Chen", "ClearPoint Clinical Solutions", "VND-004 — $280K MSA")
    
    # Spacer + RFP Section Header
    sp2 = tmp.add_paragraph()
    sp2.paragraph_format.space_before = Pt(16)
    sp2.paragraph_format.space_after = Pt(4)
    
    p_rfp_hdr = tmp.add_paragraph()
    p_rfp_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_rfp_hdr.add_run("── VENDOR CANDIDATES / RFP PROPOSALS (managed by Marcus Webb — Procurement) ──")
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string("E65100")
    
    # RFP sub-header
    p_rfp1 = tmp.add_paragraph()
    p_rfp1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_rfp1.add_run("RFP-001: Telehealth Platform")
    run.bold = True
    run.font.size = Pt(9)
    
    tr1 = tmp.add_table(rows=1, cols=3)
    tr1.alignment = WD_TABLE_ALIGNMENT.CENTER
    make_rfp_cell(tr1.cell(0, 0), "Dana Mitchell", "VirtuCare Health", "$185K/yr | Score: 7.80 ✅")
    make_rfp_cell(tr1.cell(0, 1), "Ryan Patel", "TeleMedix Solutions", "$210K/yr | Score: 8.35 ✅")
    make_rfp_cell(tr1.cell(0, 2), "Sofia Reyes", "HealthBridge Connect", "$125K/yr | Score: 6.45 ❌")
    
    p_rfp2 = tmp.add_paragraph()
    p_rfp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_rfp2.add_run("RFP-002: Clinical Analytics & Reporting")
    run.bold = True
    run.font.size = Pt(9)
    
    tr2 = tmp.add_table(rows=1, cols=2)
    tr2.alignment = WD_TABLE_ALIGNMENT.CENTER
    make_rfp_cell(tr2.cell(0, 0), "Alex Thornton", "Meridian Data Sciences", "$165K/yr | Score: 8.10 ✅")
    make_rfp_cell(tr2.cell(0, 1), "Priya Kapoor", "InsightHealth Analytics", "$130K/yr | Score: 6.90 ❌")
    
    p_rfp3 = tmp.add_paragraph()
    p_rfp3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_rfp3.add_run("RFP-003: Medical Waste Management")
    run.bold = True
    run.font.size = Pt(9)
    
    tr3 = tmp.add_table(rows=1, cols=2)
    tr3.alignment = WD_TABLE_ALIGNMENT.CENTER
    make_rfp_cell(tr3.cell(0, 0), "Carlos Mendez", "EnviroMed Disposal", "$82K/yr | Score: 7.70 ✅")
    make_rfp_cell(tr3.cell(0, 1), "Laura Kim", "GreenHealth Waste Services", "$68K/yr | Score: 6.85 ❌")
    
    # Add RBAC legend
    sp3 = tmp.add_paragraph()
    sp3.paragraph_format.space_before = Pt(14)
    
    p_legend = tmp.add_paragraph()
    p_legend.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_legend.add_run("RBAC Legend:  ")
    run.bold = True
    run.font.size = Pt(8)
    run = p_legend.add_run("🔓 = Has access    🚫 = Restricted    ")
    run.font.size = Pt(8)
    run = p_legend.add_run("Dark blue = Leadership    Medium blue = Manager    Light blue = Analyst/Specialist")
    run.font.size = Pt(8)
    
    # Final spacer
    tmp.add_paragraph("")
    
    # Now transplant all elements from tmp into the real doc after `intro`
    insert_after = intro
    for elem in list(tmp.element.body):
        body.insert(body.index(insert_after) + 1, elem)
        insert_after = elem
    
    return doc


def main():
    doc = Document(str(DOC_PATH))
    doc = build_hierarchy(doc)
    doc.save(str(OUT_PATH))
    print(f"✅ Saved: {OUT_PATH}")


if __name__ == "__main__":
    main()
